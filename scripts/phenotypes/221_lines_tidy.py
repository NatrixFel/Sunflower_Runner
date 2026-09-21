"""
Phase 8a: update tidy phenotypes from the breeder’s round-2 replies.

What it does:
1. Replaces emergence-to-flowering and central-to-side flowering data from the NEW
   file (a formula error in the 2023 calculation was fixed).
2. Replaces oil content, hull content, and 1000-seed weight data from the NEW file
   (2024 hull-content data were added).
3. Applies point outlier corrections from the breeder:
   - LI41/2021/height: 11 → 112
   - LI49/2021/hull content: 45 → 25
   - LI36/2023/1000-seed weight: 65 → 59
   - VA761A × LI19/2022/diameter: 29.75 → 20.9

Output: phase8_output/221_lines_tidy.parquet, hybrids_tidy_v2.parquet
"""

import sys
from pathlib import Path as _PathBoot
_p = _PathBoot(__file__).resolve().parent
while _p != _p.parent and not (_p / "paths.py").exists():
    _p = _p.parent
if str(_p) not in sys.path:
    sys.path.insert(0, str(_p))
from paths import DEPOSIT, WORK, CALLING, OUT, INTERMEDIATE, MODELS, use_models, work, calling, en_table, en_ids, en_id, is_line_id, ru_col
use_models()
from pathlib import Path
import re
import pandas as pd
import numpy as np
from openpyxl import load_workbook

ROOT = work()  # undeposited working tree
NEW  = ROOT / "6_новые данные_ответы + исходники +vcf"
OUT  = ROOT / "phase8_output"
OUT.mkdir(exist_ok=True)

# Old tidy
lines_old   = en_table(pd.read_parquet(ROOT / "phase1_output" / "01_lines_tidy.parquet"))
hybrids_old = en_table(pd.read_parquet(ROOT / "phase1_output" / "01_hybrids_tidy.parquet"))
print(f"Old: lines={len(lines_old)}, hybrids={len(hybrids_old)}")


# === A. Parse the new emergence-to-flowering file ===
fp = NEW / "линии_всходы-цветение 4 года — копия.xlsx"
wb = load_workbook(fp, data_only=True)
ws = wb.active  # first sheet
rows_phen = []
data = list(ws.iter_rows(values_only=True))
# Header: row 2 (genotype column, then 2021–2024 emergence-to-flowering, then 2021–2024 central-to-side)
# Data from row 3
header = data[1]  # 0: genotype, 1..4: emergence-to-flowering, 5..8: central-to-side
years_all = [2021, 2022, 2023, 2024]
for row in data[2:]:
    if row[0] is None:
        continue
    geno = en_id(str(row[0]).strip())
    if not is_line_id(geno):
        continue
    for i, y in enumerate(years_all):
        v = row[1 + i]
        if v is not None and pd.notna(v):
            rows_phen.append({"year": y, "genotype": geno,
                              "trait": "days_emergence_flowering",
                              "value": float(v), "replicate": np.nan,
                              "source": "lines_summary_v2"})
        v2 = row[5 + i] if len(row) > 5+i else None
        if v2 is not None and pd.notna(v2):
            rows_phen.append({"year": y, "genotype": geno,
                              "trait": "days_central_to_side",
                              "value": float(v2), "replicate": np.nan,
                              "source": "lines_summary_v2"})

new_phen = pd.DataFrame(rows_phen)
print(f"\nFrom the new 'emergence-to-flowering': {len(new_phen)} rows")
print(new_phen.groupby(["trait", "year"])["value"].agg(["count", "mean"]).round(2))


# === B. Parse the new oil / hull / 1000-seed-weight file ===
fp = NEW / "линии_масл-ть_ лузж-ть_ м1000шт 4 года — копия.xlsx"
wb = load_workbook(fp, data_only=True)
new_mat = []
# Sheet 'масл-ть': genotype + years 21,22,23,24 (R3..R56)
ws = wb["масл-ть"]
data = list(ws.iter_rows(values_only=True))
# Find the header row that names the genotype column
hdr_idx = next(i for i, r in enumerate(data) if r and any(str(c) == ru_col("genotype") for c in r if c is not None))
for row in data[hdr_idx+1:]:
    # First non-empty value should be a line ID
    cells = [c for c in row if c is not None]
    if not cells: continue
    geno_cell = next((c for c in row if c is not None and isinstance(c, str) and is_line_id(str(c))), None)
    if not geno_cell: continue
    geno = en_id(str(geno_cell).strip())
    # Values after the genotype column
    geno_idx = row.index(geno_cell)
    vals = row[geno_idx+1:geno_idx+5]
    for i, y in enumerate([2021, 2022, 2023, 2024]):
        v = vals[i] if i < len(vals) else None
        if v is not None and pd.notna(v):
            try:
                new_mat.append({"year": y, "genotype": geno, "trait": "oil_content",
                                "value": float(v), "replicate": np.nan,
                                "source": "lines_summary_v2"})
            except (ValueError, TypeError):
                pass

# Sheet 'лузж-ть': first column = year, the rest = line IDs
ws = wb["лузж-ть"]
data = list(ws.iter_rows(values_only=True))
# Find the row with line-ID headers
hdr_idx = next(i for i, r in enumerate(data) if r and any(isinstance(c, str) and is_line_id(str(c)) for c in r if c is not None))
li_names = []
for c in data[hdr_idx]:
    if c is not None and isinstance(c, str) and is_line_id(str(c)):
        li_names.append(en_id(str(c).strip()))
# Data below: first column is year, the rest are values
for row in data[hdr_idx+1:]:
    if not row or row[0] is None: continue
    try: y = int(row[0])
    except (ValueError, TypeError): continue
    if y < 2020 or y > 2030: continue
    li_idx = 0
    for c in row[1:]:
        if li_idx >= len(li_names): break
        if c is not None and pd.notna(c):
            try:
                new_mat.append({"year": y, "genotype": li_names[li_idx], "trait": "hull_content",
                                "value": float(c), "replicate": np.nan,
                                "source": "lines_summary_v2"})
            except (ValueError, TypeError): pass
        li_idx += 1

# Sheet 'м1000шт': same layout, first column is year, the rest are line IDs
ws = wb["м1000шт"]
data = list(ws.iter_rows(values_only=True))
hdr_idx = next(i for i, r in enumerate(data) if r and any(isinstance(c, str) and is_line_id(str(c)) for c in r if c is not None))
li_names = []
for c in data[hdr_idx]:
    if c is not None and isinstance(c, str) and is_line_id(str(c)):
        li_names.append(en_id(str(c).strip()))
for row in data[hdr_idx+1:]:
    if not row: continue
    # Find the year (an integer at the start of the row)
    y = None
    for c in row[:3]:
        if c is None: continue
        try:
            y_try = int(c)
            if 2020 <= y_try <= 2030:
                y = y_try; break
        except (ValueError, TypeError): pass
    if y is None: continue
    # Values after the year, as many as there are line IDs
    year_idx = row.index(y)
    vals = row[year_idx+1:year_idx+1+len(li_names)]
    for i, li in enumerate(li_names):
        if i >= len(vals): break
        v = vals[i]
        if v is not None and pd.notna(v):
            try:
                new_mat.append({"year": y, "genotype": li, "trait": "seed_weight_1000",
                                "value": float(v), "replicate": np.nan,
                                "source": "lines_summary_v2"})
            except (ValueError, TypeError): pass

new_mat_df = pd.DataFrame(new_mat)
print(f"\nFrom the new 'oil/hull/1000-seed-weight': {len(new_mat_df)} rows")
print(new_mat_df.groupby(["trait", "year"])["value"].agg(["count", "mean"]).round(2))


# === C. Merge: replace the affected traits ===
TRAITS_TO_REPLACE = {"days_emergence_flowering", "days_central_to_side",
                     "oil_content", "hull_content", "seed_weight_1000"}
lines_v2 = lines_old[~lines_old["trait"].isin(TRAITS_TO_REPLACE)].copy()
lines_v2 = pd.concat([lines_v2, new_phen, new_mat_df], ignore_index=True)

# Fill any missing columns to match the original
for c in lines_old.columns:
    if c not in lines_v2.columns:
        lines_v2[c] = np.nan
lines_v2 = lines_v2[lines_old.columns]
lines_v2["confectionery"] = lines_v2["genotype"].isin({"LI29", "LI30"})
print(f"\nLines v2: {len(lines_v2)} rows")


# === D. Point corrections (round 2) ===
fixes = [
    {"year": 2021, "genotype": "LI41", "trait": "plant_height",        "old": 11.0,   "new": 112.0},
    {"year": 2021, "genotype": "LI49", "trait": "hull_content",   "old": 45.0,   "new": 25.0},
    {"year": 2023, "genotype": "LI36", "trait": "seed_weight_1000",     "old": 65.0,   "new": 59.0},
]
for fix in fixes:
    mask = ((lines_v2["year"] == fix["year"]) &
            (lines_v2["genotype"] == fix["genotype"]) &
            (lines_v2["trait"] == fix["trait"]) &
            (abs(lines_v2["value"] - fix["old"]) < 0.1))
    if mask.sum() > 0:
        lines_v2.loc[mask, "value"] = fix["new"]
        print(f"  Corrected: {fix['genotype']}/{fix['year']}/{fix['trait']}: {fix['old']} → {fix['new']} ({mask.sum()} rows)")
    else:
        print(f"  NOT FOUND for correction: {fix['genotype']}/{fix['year']}/{fix['trait']}={fix['old']}")

# === E. Hybrids — correct VA761A × LI19 / 2022 / head diameter 29.75 → 20.9 ===
hyb_v2 = hybrids_old.copy()
mask = ((hyb_v2["mother_long"] == "VA761A") & (hyb_v2["father"] == "LI19") &
        (hyb_v2["year"] == 2022) & (hyb_v2["trait"] == "head_diameter") &
        (abs(hyb_v2["value"] - 29.75) < 0.1))
if mask.sum() > 0:
    hyb_v2.loc[mask, "value"] = 20.9
    print(f"\nCorrected: VA761A × LI19/2022/head_diameter: 29.75 → 20.9 ({mask.sum()} rows)")
else:
    print(f"\nNOT FOUND VA761A × LI19/2022 head_diameter 29.75")


# Save
lines_v2.to_parquet(OUT / "221_lines_tidy.parquet", index=False)
hyb_v2.to_parquet(OUT / "hybrids_tidy_v2.parquet", index=False)
print(f"\nSaved: {OUT}/221_lines_tidy.parquet and hybrids_tidy_v2.parquet")
print(f"  lines v2: {len(lines_v2)} rows, {lines_v2['trait'].nunique()} traits")
print(f"  hybrids v2: {len(hyb_v2)} rows")
